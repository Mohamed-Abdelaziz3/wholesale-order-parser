"""Product catalog loading, parsing, and validation.

The catalog is the highest-friction part of onboarding a real merchant: they
hand over a file exported from Excel, in an unpredictable encoding, with Arabic
or English headers, and possibly no alias column at all. This module absorbs
that mess so the rest of the system can assume a clean list.
"""

from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Iterable, Sequence

from .models import CatalogImportDiagnostics, CatalogProduct, CatalogRowDiagnostic
from .money import parse_money

DEFAULT_CATALOG_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "catalog.csv",
)

MAX_CATALOG_ROWS = 50_000
MAX_ALIASES_PER_PRODUCT = 40
MAX_DIAGNOSTIC_EXAMPLES = 25

# Encodings an Egyptian merchant's Excel export realistically arrives in.
# Excel's "Unicode Text (*.txt)" export is UTF-16LE with a BOM; Arabic Windows
# Excel writes cp1256. Both are what a merchant actually sends.
CANDIDATE_ENCODINGS = ("utf-8-sig", "utf-8", "utf-16", "cp1256", "latin-1")

# Accepted header spellings, normalised to lowercase/stripped.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "product_id": (
        "product_id", "productid", "id", "sku", "code", "product code",
        "كود", "الكود", "كود المنتج", "كود الصنف", "رقم الصنف", "رقم المنتج",
    ),
    "product_name": (
        "product_name", "productname", "name", "product", "description",
        "اسم المنتج", "الاسم", "المنتج", "الصنف", "اسم الصنف", "البيان",
    ),
    "aliases": (
        "aliases", "alias", "synonyms", "keywords", "other names",
        "مرادفات", "المرادفات", "اسماء بديلة", "أسماء بديلة", "اسماء اخرى", "أسماء أخرى",
    ),
    "unit": (
        "unit", "uom", "unit of measure", "packaging",
        "الوحدة", "الوحده", "وحدة", "وحده", "وحدة القياس", "التعبئة",
    ),
    "price": (
        "price", "unit price", "cost", "amount",
        "السعر", "سعر", "سعر الوحدة", "سعر البيع", "الثمن",
    ),
}

ALIAS_SEPARATORS = re.compile(r"[|;،]")
_ARABIC_INDIC = str.maketrans("٠١٢٣٤٥٦٧٨٩٫", "0123456789.")


class CatalogError(ValueError):
    """Raised when an uploaded catalog cannot be accepted safely.

    The normal error string is readable in a browser or CLI. The structured
    fields give the API enough evidence to show exactly which source rows need
    repair without installing a partial catalog.
    """

    def __init__(
        self,
        message: str,
        *,
        diagnostics: CatalogImportDiagnostics | None = None,
        detected_column_mapping: dict[str, str] | None = None,
        unmapped_source_columns: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.diagnostics = diagnostics or CatalogImportDiagnostics()
        self.detected_column_mapping = detected_column_mapping or {}
        self.unmapped_source_columns = unmapped_source_columns or []


@dataclass
class CatalogParseResult:
    """Outcome of parsing an uploaded catalog file."""

    products: list[CatalogProduct] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    detected_encoding: str = ""
    source_format: str = "csv"
    detected_column_mapping: dict[str, str] = field(default_factory=dict)
    unmapped_source_columns: list[str] = field(default_factory=list)
    diagnostics: CatalogImportDiagnostics = field(default_factory=CatalogImportDiagnostics)

    @property
    def count(self) -> int:
        return len(self.products)


def _normalise_header(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("﻿", "").strip()).lower()


def _resolve_columns(headers: Sequence[str]) -> dict[str, str]:
    """Map canonical field names to the actual header strings present."""
    normalised = {_normalise_header(h): h for h in headers if h is not None}
    resolved: dict[str, str] = {}
    for canonical, options in COLUMN_ALIASES.items():
        for option in options:
            if option in normalised:
                resolved[canonical] = normalised[option]
                break
    return resolved


def _unmapped_columns(headers: Sequence[str], columns: dict[str, str]) -> list[str]:
    """Return source headers deliberately not used by this importer.

    Unknown columns are not guessed at or silently mapped. They are harmless
    for a merchant export, but surfaced so an operator can verify the detected
    mapping before relying on the newly installed catalog.
    """
    selected = set(columns.values())
    return [header for header in headers if header and header not in selected]


def _row_example(
    diagnostics: CatalogImportDiagnostics,
    *,
    line_number: int,
    issues: list[str],
    product_id: str | None,
) -> None:
    """Add an operator example without letting a bad upload flood a response."""
    if len(diagnostics.malformed_row_examples) >= MAX_DIAGNOSTIC_EXAMPLES:
        return
    diagnostics.malformed_row_examples.append(
        CatalogRowDiagnostic(
            line_number=line_number,
            product_id=product_id or None,
            issues=issues,
        )
    )


def _validation_error_message(diagnostics: CatalogImportDiagnostics) -> str:
    """Summarise every blocking condition in a concise bilingual message."""
    facts: list[str] = []
    if diagnostics.row_limit_exceeded:
        facts.append(f"row limit exceeded / تجاوز الحد الأقصى {MAX_CATALOG_ROWS} صف")
    if diagnostics.malformed_row_count:
        facts.append(
            f"malformed rows / صفوف غير صالحة: {diagnostics.malformed_row_count}"
        )
    if diagnostics.duplicate_sku_count:
        facts.append(
            f"duplicate SKU rows / أكواد مكررة: {diagnostics.duplicate_sku_count}"
        )
    if diagnostics.missing_sku_count:
        facts.append(f"missing SKU / كود مفقود: {diagnostics.missing_sku_count}")
    if diagnostics.missing_name_count:
        facts.append(f"missing name / اسم مفقود: {diagnostics.missing_name_count}")
    if diagnostics.missing_price_count:
        facts.append(f"missing price / سعر مفقود: {diagnostics.missing_price_count}")
    if diagnostics.invalid_price_count:
        facts.append(
            f"invalid price / سعر غير صالح: {diagnostics.invalid_price_count}"
        )
    if diagnostics.missing_unit_count:
        facts.append(f"missing unit / وحدة مفقودة: {diagnostics.missing_unit_count}")
    summary = "; ".join(facts) or "catalog validation failed / فشل التحقق من الكتالوج"
    return (
        "Catalog was not installed; fix the listed source rows and upload again. "
        "لم يتم استبدال الكتالوج؛ صحح الصفوف المذكورة ثم ارفع الملف مرة أخرى. "
        f"Details: {summary}."
    )


def _has_arabic(text: str) -> bool:
    return any("؀" <= ch <= "ۿ" for ch in text)


def _decode(raw: bytes) -> tuple[str, str]:
    """Decode bytes, returning ``(text, encoding_used)``."""
    for encoding in CANDIDATE_ENCODINGS:
        try:
            text = raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        # cp1256 will "successfully" decode UTF-8 Arabic into mojibake, so it is
        # only accepted when the result actually contains Arabic. latin-1 is the
        # last-resort fallback and is deliberately NOT gated that way: it can
        # never produce an Arabic codepoint, so gating it made the branch dead
        # code and left accented Latin catalogs unreadable.
        if encoding == "cp1256" and not _has_arabic(text):
            continue
        if encoding == "utf-16" and "\x00" in text:
            continue
        return text, encoding
    return raw.decode("utf-8", errors="replace"), "utf-8 (with replacements)"


class _CatalogPriceError(ValueError):
    """A price-cell failure classified for catalog diagnostics."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


_PRICE_PATTERN = re.compile(r"^(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?$")


def _parse_catalog_price(value: object) -> Decimal:
    """Read one explicit catalog price without guessing or repairing it.

    Currency symbols, loose text, excess decimals, negative numbers, and
    malformed grouping are all rejected. Arabic digits and the Arabic decimal
    and thousands separators are deterministic spelling equivalents, so they
    are normalised before validation. A numeric zero remains valid: "free" is
    a commercial decision, whereas a missing price is not.
    """
    if value is None:
        raise _CatalogPriceError("missing_price", "price is blank")
    if isinstance(value, bool):
        raise _CatalogPriceError("invalid_price", "price is not a number")

    text = str(value).translate(_ARABIC_INDIC).strip()
    if not text:
        raise _CatalogPriceError("missing_price", "price is blank")
    text = text.replace("\u066b", ".").replace("\u066c", ",")
    if text.startswith("-"):
        raise _CatalogPriceError("invalid_price", "price must not be negative")
    if not _PRICE_PATTERN.fullmatch(text):
        raise _CatalogPriceError(
            "invalid_price",
            "price must be a non-negative number with at most two decimal places",
        )
    try:
        # Thousands separators have already been proven structurally valid.
        return parse_money(text.replace(",", ""), "Catalog price")
    except ValueError as exc:
        raise _CatalogPriceError("invalid_price", str(exc)) from None


def _split_aliases(value: object) -> list[str]:
    if not value:
        return []
    parts = ALIAS_SEPARATORS.split(str(value))
    seen: list[str] = []
    for part in parts:
        cleaned = re.sub(r"\s+", " ", part).strip()
        if cleaned and cleaned not in seen:
            seen.append(cleaned)
    return seen[:MAX_ALIASES_PER_PRODUCT]


def _rows_from_xlsx(raw: bytes) -> list[list[object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise CatalogError(
            "ملفات Excel تحتاج مكتبة openpyxl. ثبّتها أو صدّر الملف بصيغة CSV."
        ) from exc
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _looks_like_xlsx(raw: bytes, filename: str) -> bool:
    return filename.lower().endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK"


def parse_catalog_bytes(raw: bytes, filename: str = "catalog.csv") -> CatalogParseResult:
    """Parse an uploaded catalog file into validated products.

    Accepts CSV (several encodings, comma/semicolon/tab delimited) and XLSX.
    Raises :class:`CatalogError` with an operator-readable Arabic message when
    the file cannot be used.
    """
    if not raw or not raw.strip():
        raise CatalogError("الملف فارغ.")

    result = CatalogParseResult()

    if _looks_like_xlsx(raw, filename):
        raw_rows = _rows_from_xlsx(raw)
        result.source_format = "xlsx"
        result.detected_encoding = "xlsx"
        if not raw_rows:
            raise CatalogError("الملف لا يحتوي على أي صفوف.")
        headers = [str(cell) if cell is not None else "" for cell in raw_rows[0]]
        body: Iterable[dict[str, object]] = [
            {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
            for row in raw_rows[1:]
        ]
    else:
        text, encoding = _decode(raw)
        result.detected_encoding = encoding
        sample = text[:8192]
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
                sample, delimiters=",;\t"
            )
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = list(reader.fieldnames or [])
        body = reader

    columns = _resolve_columns(headers)
    unmapped_columns = _unmapped_columns(headers, columns)
    result.detected_column_mapping = dict(columns)
    result.unmapped_source_columns = unmapped_columns

    # Price is commercial authority from the catalog, so it is just as
    # required as the SKU, product name, and unit. Never manufacture a zero
    # price because a merchant omitted a column.
    required_columns = ("product_id", "product_name", "unit", "price")
    missing = [field for field in required_columns if field not in columns]
    if missing:
        result.diagnostics.missing_required_columns = list(missing)
        present = ", ".join(header for header in headers if header) or "(no headers / بدون عناوين)"
        missing_names = ", ".join(missing)
        raise CatalogError(
            "Missing required catalog column(s) / أعمدة كتالوج مطلوبة مفقودة: "
            f"{missing_names}. Detected source columns / الأعمدة الموجودة: {present}. "
            "Required aliases include SKU: product_id / sku / كود المنتج; "
            "name: product_name / اسم المنتج / الصنف; unit: unit / uom / الوحدة; "
            "price: price / unit price / السعر.",
            diagnostics=result.diagnostics,
            detected_column_mapping=columns,
            unmapped_source_columns=unmapped_columns,
        )

    products: list[CatalogProduct] = []
    seen_ids: dict[str, int] = {}
    diagnostics = result.diagnostics

    for line_number, row in enumerate(body, start=2):
        # Empty spreadsheet/CSV rows are explicitly reported, but do not turn
        # an otherwise valid merchant export into a failure.
        if not any(str(value).strip() for value in row.values() if value is not None):
            diagnostics.blank_row_count += 1
            continue

        diagnostics.data_row_count += 1
        if diagnostics.data_row_count > MAX_CATALOG_ROWS:
            diagnostics.row_limit_exceeded = True
            _row_example(
                diagnostics,
                line_number=line_number,
                product_id=None,
                issues=[f"row limit exceeded ({MAX_CATALOG_ROWS})"],
            )
            break

        product_id = str(row.get(columns["product_id"]) or "").strip()
        product_name = re.sub(
            r"\s+", " ", str(row.get(columns["product_name"]) or "")
        ).strip()
        unit = str(row.get(columns["unit"]) or "").strip()
        issues: list[str] = []

        if not product_id:
            diagnostics.missing_sku_count += 1
            issues.append("missing SKU / كود المنتج مفقود")
        elif product_id in seen_ids:
            diagnostics.duplicate_sku_count += 1
            issues.append(
                f"duplicate SKU '{product_id}' (first seen on row {seen_ids[product_id]})"
            )
        else:
            seen_ids[product_id] = line_number

        if not product_name:
            diagnostics.missing_name_count += 1
            issues.append("missing product name / اسم المنتج مفقود")
        if not unit:
            diagnostics.missing_unit_count += 1
            issues.append("missing unit / وحدة الصنف مفقودة")

        price: Decimal | None = None
        try:
            price = _parse_catalog_price(row.get(columns["price"]))
        except _CatalogPriceError as exc:
            if exc.code == "missing_price":
                diagnostics.missing_price_count += 1
                issues.append("missing price / سعر الصنف مفقود")
            else:
                diagnostics.invalid_price_count += 1
                issues.append(f"invalid price / سعر غير صالح: {exc}")

        if issues:
            diagnostics.malformed_row_count += 1
            _row_example(
                diagnostics,
                line_number=line_number,
                product_id=product_id or None,
                issues=issues,
            )
            continue

        aliases = _split_aliases(row.get(columns["aliases"])) if "aliases" in columns else []
        # ``price`` is present exactly when no price issue was recorded.
        assert price is not None
        products.append(
            CatalogProduct(
                product_id=product_id,
                product_name=product_name,
                aliases=aliases,
                unit=unit,
                price=price,
            )
        )

    if diagnostics.row_limit_exceeded or diagnostics.malformed_row_count:
        raise CatalogError(
            _validation_error_message(diagnostics),
            diagnostics=diagnostics,
            detected_column_mapping=columns,
            unmapped_source_columns=unmapped_columns,
        )

    if not products:
        raise CatalogError(
            "No usable catalog products were found / لم يتم العثور على أصناف صالحة في الملف.",
            diagnostics=diagnostics,
            detected_column_mapping=columns,
            unmapped_source_columns=unmapped_columns,
        )

    if diagnostics.blank_row_count:
        result.warnings.append(
            f"Blank rows skipped / تم تجاهل صفوف فارغة: {diagnostics.blank_row_count}."
        )
    if "aliases" not in columns:
        result.warnings.append(
            "No aliases column / لا يوجد عمود مرادفات. إضافة أسماء بالعامية لكل صنف "
            "ترفع دقة المطابقة بشكل ملحوظ."
        )
    if unmapped_columns:
        result.warnings.append(
            "Unmapped source columns / أعمدة لم تُستخدم: " + ", ".join(unmapped_columns)
        )

    result.products = products
    return result


def load_catalog(csv_path: str | None = None) -> list[CatalogProduct]:
    """Load a catalog from disk (bundled demo catalog and first-run seeding)."""
    if csv_path is None:
        csv_path = os.getenv("CATALOG_PATH") or DEFAULT_CATALOG_PATH

    path = Path(csv_path)
    if not path.exists():
        raise FileNotFoundError(f"Catalog file not found: {csv_path}")

    return parse_catalog_bytes(path.read_bytes(), path.name).products


def get_catalog_dict(products: list[CatalogProduct]) -> dict[str, CatalogProduct]:
    """Create a ``product_id -> CatalogProduct`` lookup dict."""
    return {p.product_id: p for p in products}


def catalog_template_csv() -> str:
    """A ready-to-fill template the merchant can export their catalog into."""
    buffer = io.StringIO()
    buffer.write("﻿")
    writer = csv.writer(buffer)
    writer.writerow(["product_id", "product_name", "aliases", "unit", "price"])
    writer.writerow(
        [
            "CL001",
            "صابون سائل لليدين كبير 2 لتر",
            "صابون سايل كبير|صابون ايد كبير",
            "قطعة",
            "45.00",
        ]
    )
    writer.writerow(["CL004", "كلور مركز 4 لتر", "كلوركس مركز|كلور كبير", "جركن", "35.00"])
    return buffer.getvalue()
