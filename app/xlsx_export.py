"""The formatted ``.xlsx`` an operator hands to a distributor owner.

Why this is a separate module
-----------------------------
The CSV builder is inlined into its route in :mod:`app.main`, which is a large
part of why that file is as long as it is. The workbook is a good deal more
code than the CSV, so the shaping lives here and the route stays short. This
mirrors :mod:`app.documents`, which holds the picking note's shaping outside
the request handler.

What the file may and may not claim
-----------------------------------
Only an approved snapshot ever reaches this module, and every line in an
approved snapshot was confirmed by a human — ``app.routing`` disables
auto-accept unconditionally. ``confidence`` is therefore the matcher's
*advisory* score and nothing else: it is kept under the same
``Model confidence (advisory)`` header the CSV uses, and the colour scale is
explained by a legend under the totals. A red cell means the model was unsure,
never that the line is unapproved.

Every string written here is untrusted — product names come from
merchant-uploaded catalogs and customer WhatsApp text. openpyxl does not escape
formulas, so Excel would evaluate a cell whose text begins with ``=``. Every
string cell goes through :func:`app.security.csv_safe`; numbers are written as
real numbers so the merchant can still sum a column.
"""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .documents import document_date

# The thresholds the matcher itself routes on, so the colours in the file mean
# what they mean everywhere else in the system.
from .matcher import CONFIDENCE_HIGH, CONFIDENCE_MEDIUM
from .security import csv_safe

# Column headers, verbatim from the CSV export, so the two files stay
# recognisably the same document.
COLUMNS = (
    "#",
    "كود المنتج",
    "اسم المنتج",
    "الكمية",
    "الوحدة",
    "السعر الفردي",
    "الإجمالي",
    "Model confidence (advisory)",
    "Human confirmed",
    "Human actor",
    "تعديل يدوي (Corrected)",
    "Catalog unit price",
    "Price override",
)

CONFIDENCE_COLUMN = 8  # 1-based index of "Model confidence (advisory)"

# The totals sit under the column they sum, as they do in the CSV, with the
# label in the cell beside them rather than stranded at the far end of the row.
TOTAL_LABEL_COLUMN = 6
TOTAL_VALUE_COLUMN = 7

# xlsx references fonts by name, not by file. The client's Excel will not have
# the UI's Cairo installed, so this is a name that ships with Windows and has
# Arabic coverage; Excel falls back sanely elsewhere.
FONT_NAME = "Segoe UI"

MONEY_FORMAT = "#,##0.00"
QUANTITY_FORMAT = "#,##0.###"
PERCENT_FORMAT = "0%"

HEADER_FILL = PatternFill(start_color="FF1F3864", end_color="FF1F3864", fill_type="solid")
LABEL_FILL = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")

# Excel's own Good / Neutral / Bad pairs — familiar to anyone who has used
# conditional formatting, and legible when printed in greyscale.
CONFIDENCE_STYLES = (
    (CONFIDENCE_HIGH, "FFC6EFCE", "FF006100"),
    (CONFIDENCE_MEDIUM, "FFFFEB9C", "FF9C6500"),
    (None, "FFFFC7CE", "FF9C0006"),
)

THIN = Side(style="thin", color="FFBFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Per-column horizontal alignment. rightToLeft flips the sheet's reading
# direction but does not set cell alignment, and Google Sheets ignores the
# sheet-level flag for cell content entirely, so this is set explicitly.
ALIGNMENTS = (
    "center",  # #
    "center",  # كود المنتج
    "right",  # اسم المنتج
    "center",  # الكمية
    "right",  # الوحدة
    "center",  # السعر الفردي
    "center",  # الإجمالي
    "center",  # Model confidence
    "center",  # Human confirmed
    "right",  # Human actor
    "right",  # تعديل يدوي
    "center",  # Catalog unit price
    "center",  # Price override
)

WIDTH_MIN = 8
WIDTH_MAX = 50
PRODUCT_NAME_FLOOR = 30
LEGEND = (
    "النسبة دي تقدير آلي استرشادي فقط (Model confidence). كل سطر في الملف ده "
    "اتأكد من موظف بالاسم — اللون الأحمر معناه إن النموذج مكانش متأكد، "
    "مش إن السطر مش معتمد."
)


def _default_title(settings: Dict[str, Any]) -> str:
    """The merchant's letterhead line: their shop, else the note title."""
    shop_name = str(settings.get("shop_name") or "").strip()
    document_title = str(settings.get("document_title") or "").strip()
    return shop_name or document_title or "طلب معتمد"


def filename_stem(export_data: Dict[str, Any], settings: Dict[str, Any], order_id: int) -> str:
    """Name the file after the merchant, as the document route does."""
    shop_name = str(settings.get("shop_name") or "").strip()
    customer_name = str(export_data.get("customer_name") or "").strip()
    return shop_name or customer_name or f"order-{order_id}"


def export_date(export_data: Dict[str, Any]) -> str:
    """``YYYY-MM-DD`` for the filename; the sheet shows ``YYYY/MM/DD``."""
    return document_date(export_data.get("last_exported_at")).replace("/", "-")


def _number(value: Any) -> Optional[float]:
    """A real number for Excel, or ``None`` so the cell stays empty."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _confidence_style(confidence: float) -> tuple[str, str]:
    for threshold, fill, colour in CONFIDENCE_STYLES:
        if threshold is None or confidence >= threshold:
            return fill, colour
    return CONFIDENCE_STYLES[-1][1], CONFIDENCE_STYLES[-1][2]


def _write(ws, row: int, column: int, value: Any, **style) -> Any:
    """Write one cell with every string neutralised against formula injection."""
    cell = ws.cell(row=row, column=column, value=csv_safe(value))
    cell.font = style.get("font") or Font(name=FONT_NAME)
    if style.get("alignment") is not None:
        cell.alignment = style["alignment"]
    if style.get("fill") is not None:
        cell.fill = style["fill"]
    if style.get("border") is not None:
        cell.border = style["border"]
    if style.get("number_format"):
        cell.number_format = style["number_format"]
    return cell


def _label_row(ws, row: int, label: str, value: Any, *, wrap: bool = False) -> None:
    """One ``label | value`` pair, the value merged across the table width."""
    _write(
        ws,
        row,
        1,
        label,
        font=Font(name=FONT_NAME, bold=True),
        alignment=Alignment(horizontal="right", vertical="center"),
        fill=LABEL_FILL,
        border=CELL_BORDER,
    )
    _write(
        ws,
        row,
        2,
        value,
        alignment=Alignment(horizontal="right", vertical="center", wrap_text=wrap),
        border=CELL_BORDER,
    )
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(COLUMNS))


def _apply_column_widths(ws, rendered: List[List[str]]) -> None:
    """openpyxl has no autofit, so widths are measured from the content.

    Arabic glyphs are narrower than the width unit assumes; the multiplier is
    padding rather than a faithful text measurement, and the clamp is what
    stops one long product name from producing an unusable column.
    """
    for index in range(1, len(COLUMNS) + 1):
        longest = max((len(row[index - 1]) for row in rendered), default=0)
        width = longest * 1.1 + 2
        floor = PRODUCT_NAME_FLOOR if index == 3 else WIDTH_MIN
        ws.column_dimensions[get_column_letter(index)].width = max(
            floor, min(WIDTH_MAX, width)
        )


def build_order_workbook(
    export_data: Dict[str, Any],
    settings: Dict[str, Any],
    order_id: int,
) -> bytes:
    """Render one approved snapshot as a formatted single-sheet workbook.

    ``export_data`` is the dict returned by ``database.record_export`` and
    ``settings`` the dict from ``database.get_shop_settings``. Snapshot rows are
    raw SQLite rows, so the boolean columns arrive as 0/1 ints and are read for
    truthiness rather than identity.
    """
    workbook = Workbook()
    ws = workbook.active
    ws.title = "الطلب"
    ws.sheet_view.rightToLeft = True

    width = len(COLUMNS)
    last_column = get_column_letter(width)

    # Row 1 — the letterhead the merchant retypes if they want to.
    ws.merge_cells(f"A1:{last_column}1")
    _write(
        ws,
        1,
        1,
        _default_title(settings),
        font=Font(name=FONT_NAME, bold=True, size=16),
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    ws.row_dimensions[1].height = 28

    # Header block. The customer's message is stated once, here, above the
    # table — it used to be repeated on every line, which pushed the product
    # columns off the screen.
    _label_row(ws, 2, "اسم العميل", export_data.get("customer_name", ""))
    _label_row(ws, 3, "هاتف العميل", export_data.get("customer_phone", ""))
    _label_row(ws, 4, "عنوان العميل", export_data.get("customer_address", ""))
    _label_row(ws, 5, "رقم الطلب (Order ID)", order_id)
    _label_row(ws, 6, "تاريخ التصدير", document_date(export_data.get("last_exported_at")))
    _label_row(ws, 7, "رسالة العميل", export_data.get("original_text", ""), wrap=True)
    ws.row_dimensions[7].height = 48

    header_row = 9
    for index, title in enumerate(COLUMNS, start=1):
        _write(
            ws,
            header_row,
            index,
            title,
            font=Font(name=FONT_NAME, bold=True, color="FFFFFFFF"),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
            fill=HEADER_FILL,
            border=CELL_BORDER,
        )
    ws.row_dimensions[header_row].height = 32

    # Column/row based, and not mirrored by rightToLeft, so "A" is still the
    # correct anchor on an RTL sheet.
    ws.freeze_panes = f"A{header_row + 1}"

    rendered: List[List[str]] = [list(COLUMNS)]
    order_total = 0.0
    row = header_row

    for index, item in enumerate(export_data["snapshot"], start=1):
        row = header_row + index
        price = float(item.get("price", 0.0) or 0.0)
        quantity = float(item.get("quantity", 0.0) or 0.0)
        line_total = round(price * quantity, 2)
        order_total += line_total
        confidence = float(item.get("confidence", 0.0) or 0.0)

        values = [
            index,
            item.get("product_id", ""),
            item.get("product_name", ""),
            quantity,
            item.get("unit", ""),
            price,
            line_total,
            confidence,
            "yes" if item.get("is_human_confirmed") else "no",
            item.get("human_actor", ""),
            "نعم" if item.get("is_manually_corrected") else "لا",
            _number(item.get("catalog_price")),
            _number(item.get("price_override")),
        ]
        formats = {
            4: QUANTITY_FORMAT,
            6: MONEY_FORMAT,
            7: MONEY_FORMAT,
            8: PERCENT_FORMAT,
            12: MONEY_FORMAT,
            13: MONEY_FORMAT,
        }
        for column, value in enumerate(values, start=1):
            cell = _write(
                ws,
                row,
                column,
                value,
                alignment=Alignment(
                    horizontal=ALIGNMENTS[column - 1], vertical="center", wrap_text=column == 3
                ),
                border=CELL_BORDER,
                number_format=formats.get(column),
            )
            if column == CONFIDENCE_COLUMN:
                fill, colour = _confidence_style(confidence)
                cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
                cell.font = Font(name=FONT_NAME, bold=True, color=colour)

        rendered.append([str(value if value is not None else "") for value in values])

    discount = float(export_data.get("discount") or 0.0)
    try:
        grand_total = round(float(export_data.get("grand_total")), 2)
    except (TypeError, ValueError):
        grand_total = round(max(0.0, round(order_total, 2) - discount), 2)

    totals_row = row + 2
    totals = (
        ("الإجمالي قبل الخصم", round(order_total, 2), False),
        ("الخصم", discount, False),
        ("الإجمالي النهائي", grand_total, True),
    )
    for offset, (label, amount, bold) in enumerate(totals):
        current = totals_row + offset
        _write(
            ws,
            current,
            TOTAL_LABEL_COLUMN,
            label,
            font=Font(name=FONT_NAME, bold=True),
            alignment=Alignment(horizontal="right", vertical="center"),
            border=CELL_BORDER,
        )
        _write(
            ws,
            current,
            TOTAL_VALUE_COLUMN,
            amount,
            font=Font(name=FONT_NAME, bold=bold),
            alignment=Alignment(horizontal="center", vertical="center"),
            border=CELL_BORDER,
            number_format=MONEY_FORMAT,
        )

    legend_row = totals_row + len(totals) + 1
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=width)
    _write(
        ws,
        legend_row,
        1,
        LEGEND,
        font=Font(name=FONT_NAME, size=9, italic=True),
        alignment=Alignment(horizontal="right", vertical="center", wrap_text=True),
    )
    ws.row_dimensions[legend_row].height = 30

    _apply_column_widths(ws, rendered)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
